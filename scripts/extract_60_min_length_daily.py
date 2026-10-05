r"""Imperial hourly length bins -> tab-delimited Access table.

Run normally after editing SOURCE_FOLDER, or:
  python extract_60_min_length_daily.py --source "C:/Counts"
  python extract_60_min_length_daily.py --source "C:/Counts/file_60 Min Length.xlsx"

Searches all subfolders, excluding directories named VOID (case-insensitive).
Reads source files directly and never changes them. Each run has a new output
folder. Direction and the original length-bin labels are preserved; units are
not inferred. H01 is 00:00-01:00 and H24 is 23:00-24:00. Missing/invalid counts
remain blank, not zero. Total is the sum of available hourly counts only.
CompleteDay is 1 only when all 24 hourly counts in that bin are valid.
Duplicate timestamps reject the workbook; duplicate station/date/direction/bin
keys across workbooks are reported and retained for review before import.
Requires openpyxl. Suggested Access table name: tblLengthByHour.
"""
from __future__ import annotations
import argparse
import csv
import fnmatch
import os
import re
import sys
import tempfile
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

SOURCE_FOLDER = Path(r"C:\Users\Kschellinger\SJTPO_Git\TAZ_Modular\input\02_May ATRs")
OUTPUT_FOLDER = Path(r"C:\Users\Kschellinger\SJTPO_Git\TAZ_Modular\traffic_counts\processed")
PATTERN = "*60 Min Length*.xlsx"
FIELDS = ['Station', 'CountDate', 'Direction', 'Route', 'LengthBin'] + [
    f'H{h:02d}' for h in range(1, 25)
] + ['Total', 'HoursPresent', 'CompleteDay', 'SourceFile', 'SourceSheet']


def parse_time(value):
    if isinstance(value, datetime):
        return value
    for fmt in ('%m/%d/%Y %I:%M %p', '%m/%d/%Y %I:%M:%S %p',
                '%m/%d/%Y %H:%M', '%m/%d/%Y %H:%M:%S',
                '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M'):
        try:
            return datetime.strptime(str(value).strip(), fmt)
        except ValueError:
            pass
    return None


def parse_count(value):
    if value is None or str(value).strip() == '':
        return None
    try:
        n = Decimal(str(value).strip())
        if n.is_finite() and n >= 0 and n == n.to_integral_value():
            return int(n)
    except InvalidOperation:
        pass
    return None


def discover(source):
    if source.is_file():
        if source.suffix.lower() != '.xlsx' or source.name.startswith('~$'):
            raise ValueError('Select an .xlsx workbook, not an Excel temporary file.')
        if any(part.casefold() == 'void' for part in source.parts):
            raise ValueError('Selected workbook is in a VOID folder.')
        return [source]
    if not source.is_dir():
        raise FileNotFoundError(f'Source folder does not exist: {source}')
    if source.name.casefold() == 'void':
        raise ValueError('The source folder is named VOID.')
    def onerror(error):
        raise error
    paths = []
    for folder, dirs, names in os.walk(source, onerror=onerror, followlinks=False):
        dirs[:] = [d for d in dirs if d.casefold() != 'void']
        paths.extend(Path(folder) / n for n in names
                     if not n.startswith('~$') and fnmatch.fnmatchcase(n.lower(), PATTERN.lower()))
    return sorted(paths, key=lambda p: str(p).casefold())


def extract(path, label):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    output, warnings = [], []
    try:
        info = {}
        if 'Count Information' in wb.sheetnames:
            for row in wb['Count Information'].iter_rows(values_only=True):
                if row and row[0] is not None:
                    info[str(row[0]).strip().rstrip(':')] = next(
                        (str(v).strip() for v in row[1:] if v is not None and str(v).strip()), '')
        station = info.get('Count Name') or re.sub(r'_60\s*Min\s*Length.*$', '', path.stem, flags=re.I)
        station = re.sub(r'^\s*(?:SJTPO\s+)?ATR\s*\d+\s*[.\-_:]*\s*', '', station, flags=re.I).strip()
        for ws in wb:
            if ws.title == 'Count Information':
                continue
            iterator = ws.iter_rows(values_only=True)
            header = None
            for rowno, row in enumerate(iterator, 1):
                if str(row[0]).strip().lower() in ('date/time', 'date time', 'datetime'):
                    header = row
                    break
                if rowno >= 25:
                    break
            if header is None:
                warnings.append(f'Skipped sheet {ws.title}: no Date/Time header.')
                continue
            bins = [(i, str(v).strip()) for i, v in enumerate(header) if i and v is not None and str(v).strip()]
            if not bins or len({b for _, b in bins}) != len(bins):
                raise ValueError(f'{ws.title}: missing or duplicate bin labels')
            # Reject unexpected summary columns instead of interpreting them as bins.
            if any(not re.fullmatch(r'\s*>?\s*\d+(?:\.\d+)?(?:\s+to\s+\d+(?:\.\d+)?)?\s*', b, re.I) for _, b in bins):
                raise ValueError(f'{ws.title}: unrecognized length-bin headers; inspect workbook')
            direction, _, route = ws.title.partition(',')
            route = route.strip() or info.get('Location 1', '')
            days = {}
            seen = set()
            invalid = 0
            for row in iterator:
                if all(v is None or str(v).strip() == '' for v in row):
                    continue
                dt = parse_time(row[0])
                if dt is None:
                    raise ValueError(f'{ws.title}: unexpected date/time value {row[0]!r}')
                if dt.minute or dt.second or dt.microsecond:
                    raise ValueError(f'{ws.title}: timestamp is not on an hour boundary: {dt}')
                if dt in seen:
                    raise ValueError(f'{ws.title}: duplicate timestamp {dt}')
                seen.add(dt)
                day = days.setdefault(dt.date(), {b: {} for _, b in bins})
                for col, b in bins:
                    value = parse_count(row[col] if col < len(row) else None)
                    if value is None:
                        invalid += 1
                    day[b][dt.hour + 1] = value
            for date, values in sorted(days.items()):
                incomplete = False
                for _, b in bins:
                    hours = values[b]
                    present = sum(v is not None for v in hours.values())
                    incomplete |= present != 24
                    r = dict(Station=station, CountDate=date.strftime('%m/%d/%Y'),
                             Direction=direction.strip(), Route=route, LengthBin=b)
                    r.update({f'H{h:02d}': hours.get(h) for h in range(1, 25)})
                    r.update(Total=sum(v for v in hours.values() if v is not None) if present else None,
                             HoursPresent=present, CompleteDay=int(present == 24),
                             SourceFile=label, SourceSheet=ws.title)
                    output.append(r)
                if incomplete:
                    warnings.append(f'{ws.title}, {date}: incomplete day; consult HoursPresent.')
            if invalid:
                warnings.append(f'{ws.title}: {invalid} missing/invalid counts left blank.')
        if not output:
            raise ValueError('No hourly length-bin records found.')
        return output, warnings
    finally:
        wb.close()


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('--source', type=Path, default=SOURCE_FOLDER)
    p.add_argument('--output', type=Path, default=OUTPUT_FOLDER)
    args = p.parse_args()
    source = args.source.expanduser().resolve()
    paths = discover(source)
    if not paths:
        raise FileNotFoundError(f'No {PATTERN} files in {source} or its subfolders (VOID excluded).')
    rows, log, keys = [], [], {}
    failed = 0
    def report(msg):
        print(msg, flush=True)
        log.append(msg)
    report(f'Source: {source}; workbooks: {len(paths)}; VOID folders excluded.')
    report('Length units unspecified; original labels preserved. Directions remain separate.')
    for n, path in enumerate(paths, 1):
        label = path.name if source.is_file() else path.relative_to(source).as_posix()
        report(f'[{n}/{len(paths)}] {label}')
        try:
            data, warnings = extract(path, label)
        except Exception as exc:
            failed += 1
            report(f'FAILED: {label}: {exc}')
            continue
        for r in data:
            key = tuple(r[k] for k in ('Station','CountDate','Direction','LengthBin'))
            if key in keys:
                warnings.append(f'Duplicate key {key}: {keys[key]} and {label}; both retained, review before import.')
            keys[key] = label
        rows.extend(data)
        report(f'  Export rows: {len(data)}')
        for warning in warnings:
            report('WARNING: ' + warning)
    args.output.mkdir(parents=True, exist_ok=True)
    out = Path(tempfile.mkdtemp(prefix=datetime.now().strftime('length_%Y%m%d_%H%M%S_'), dir=args.output))
    if rows:
        dest = out / 'tblLengthByHour_import_daily.txt'
        with dest.open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=FIELDS, delimiter='\t', lineterminator='\n')
            writer.writeheader()
            writer.writerows(rows)
        report(f'Wrote {len(rows)} rows: {dest.resolve()}')
    report(f'Failed workbooks: {failed}; incomplete bin-day rows: {sum(not r["CompleteDay"] for r in rows)}')
    report(f'Log: {(out / "length_processing_log.txt").resolve()}')
    (out / 'length_processing_log.txt').write_text('\n'.join(log)+'\n', encoding='utf-8')
    return 1 if failed or not rows else 0


if __name__ == '__main__':
    try:
        import openpyxl
    except ModuleNotFoundError:
        print(f'Missing openpyxl. In PowerShell run: & "{sys.executable}" -m pip install openpyxl')
        raise SystemExit(1)
    try:
        raise SystemExit(main())
    except (OSError, ValueError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        raise SystemExit(1)
