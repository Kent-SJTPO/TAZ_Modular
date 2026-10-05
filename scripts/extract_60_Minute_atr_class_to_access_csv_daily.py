"""Extract Imperial 60 Minute classification workbooks recursively.

Run normally to use SOURCE_FOLDER below. Or use:
    python extract_60_Minute_atr_class_to_access_csv_daily.py --choose-folder
    python extract_60_Minute_atr_class_to_access_csv_daily.py --source "C:/Counts"
    python extract_60_Minute_atr_class_to_access_csv_daily.py --source "C:/Counts" --output "C:/Results"

Requires openpyxl in the Python environment running this script.
Reads original workbooks directly; never copies, moves, or modifies them.
Each run writes to a new dated directory beneath OUTPUT_FOLDER.
Output columns and class aggregation match the original script: all numeric
columns after column A are summed (directions are combined). Missing hours
remain zero; daily totals are not certified complete 24-hour counts.
"""

from __future__ import annotations

import csv
import re
import argparse
import fnmatch
import os
import sys
import tempfile
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ModuleNotFoundError as exc:
    if exc.name != "openpyxl":
        raise
    print("Missing package: openpyxl. Run this in PowerShell:", file=sys.stderr)
    print(f'& "{sys.executable}" -m pip install openpyxl', file=sys.stderr)
    raise SystemExit(1) from None


SOURCE_FOLDER = Path(
    r"C:\Users\Kschellinger\SJTPO_Git\TAZ_Modular\input\02_May ATRs"
)

ROOT_FOLDER = Path(
    r"C:\Users\Kschellinger\SJTPO_Git\TAZ_Modular\traffic_counts"
)

OUTPUT_FOLDER = ROOT_FOLDER / "processed"

INPUT_PATTERN = "*60 Minute*.xlsx"

CLASS_OUTPUT_NAME = "tblClassByHour_import_daily.txt"
STATION_OUTPUT_NAME = "tblTrafficStations_import_daily.txt"
LOG_OUTPUT_NAME = "processing_log.txt"

DEFAULT_COUNTY = ""
DEFAULT_COUNTY_ORDER = ""
DEFAULT_STATION_TYPE = "ATR"
DEFAULT_ROUTE = ""

CLASS_SHEETS = [
    "Motorcycles",
    "Lights",
    "Single-Unit Trucks",
    "Articulated Trucks",
    "Buses",
    "Bicycles on Road",
]


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def safe_int(value: Any) -> int:
    if value is None or value == "":
        return 0

    try:
        return int(value)
    except (TypeError, ValueError):
        pass

    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def as_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return datetime.combine(value, time.min)

    if value is None or value == "":
        return None

    text = str(value).strip()

    formats = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%m/%d/%Y %I:%M:%S %p",
        "%m/%d/%Y %I:%M %p",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue

    return None


def access_datetime(value: datetime | None) -> str:
    if value is None:
        return ""
    return value.strftime("%m/%d/%Y %I:%M:%S %p")


def access_date(value: date | None) -> str:
    if value is None:
        return ""
    return value.strftime("%m/%d/%Y")


def read_summary(wb: Any) -> dict[str, Any]:
    summary = {}

    if "Summary" not in wb.sheetnames:
        return summary

    ws = wb["Summary"]

    for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        key = clean_text(row[0])
        value = row[1]

        if key:
            summary[key] = value

    return summary


def remove_provider_atr_prefix(text: str) -> str:
    text = clean_text(text)

    text = re.sub(
        r"^\s*ATR\s*\d+\s*[\.\-_:]*\s*",
        "",
        text,
        flags=re.IGNORECASE,
    )

    text = re.sub(
        r"^\s*SJTPO\s+ATR\s*\d+\s*[\.\-_:]*\s*",
        "",
        text,
        flags=re.IGNORECASE,
    )

    return text.strip()


def derive_base_station_id(summary: dict[str, Any], source_file: Path) -> str:
    study_name = clean_text(summary.get("Study Name"))
    station = remove_provider_atr_prefix(study_name)

    if station:
        return station[:255]

    name = source_file.stem
    name = re.sub(r"_60\s*Minute$", "", name, flags=re.IGNORECASE)
    station = remove_provider_atr_prefix(name)

    if station:
        return station[:255]

    return "UNKNOWN"


def derive_description(summary: dict[str, Any], source_file: Path) -> str:
    study_name = clean_text(summary.get("Study Name"))

    if study_name:
        return study_name[:255]

    return source_file.stem[:255]


def read_hourly_by_day(ws: Any) -> tuple[dict[date, dict[int, int]], dict[date, int]]:
    hourly_by_day = defaultdict(lambda: defaultdict(int))
    records_by_day = defaultdict(int)

    for row in ws.iter_rows(min_row=4, values_only=True):
        start_dt = as_datetime(row[0])

        if start_dt is None:
            continue

        count_date = start_dt.date()
        hour_number = start_dt.hour + 1

        row_total = sum(safe_int(value) for value in row[1:])

        hourly_by_day[count_date][hour_number] += row_total
        records_by_day[count_date] += 1

    return hourly_by_day, records_by_day


def write_delimited(path: Path, fieldnames: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=fieldnames,
            delimiter="\t",
            quoting=csv.QUOTE_MINIMAL,
            lineterminator="\n",
            extrasaction="ignore",
        )

        writer.writeheader()

        for row in rows:
            writer.writerow(row)


def write_log(path: Path, lines: list[str]) -> None:
    with path.open("w", encoding="utf-8") as file:
        for line in lines:
            file.write(line + "\n")


def discover_workbooks(source: Path, pattern: str) -> list[Path]:
    """Search every level; stop on inaccessible folders instead of hiding omissions."""
    if not source.is_dir():
        raise FileNotFoundError(f"Source folder not found: {source}")

    def report_error(error: OSError) -> None:
        raise error

    files = []
    for folder, _, names in os.walk(source, onerror=report_error, followlinks=False):
        for name in names:
            if name.startswith("~$") or not name.lower().endswith(".xlsx"):
                continue
            if fnmatch.fnmatchcase(name.lower(), pattern.lower()):
                files.append(Path(folder) / name)
    return sorted(files, key=lambda path: str(path).casefold())


def choose_source(initial: Path) -> Path | None:
    """Optional Windows folder picker; cancellation exits without writing files."""
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        try:
            selected = filedialog.askdirectory(
                title="Select the parent folder containing ATR workbooks",
                initialdir=str(initial if initial.is_dir() else Path.home()),
                mustexist=True,
            )
        finally:
            root.destroy()
    except Exception as exc:
        raise RuntimeError(
            "Could not open the folder picker. Use --source or edit SOURCE_FOLDER. "
            f"Details: {exc}"
        ) from exc
    return Path(selected) if selected else None


def process_workbook(
    xlsx_path: Path,
    source_label: str | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[str], bool]:

    warnings = []
    class_rows = []
    station_rows = []
    success = False

    wb = None

    try:
        wb = load_workbook(xlsx_path, data_only=True, read_only=True)

        summary = read_summary(wb)

        base_station_id = derive_base_station_id(summary, xlsx_path)
        description = derive_description(summary, xlsx_path)
        source_file = source_label or xlsx_path.name

        all_data = defaultdict(dict)
        records_by_station_day = defaultdict(int)

        for sheet_name in CLASS_SHEETS:
            if sheet_name not in wb.sheetnames:
                warnings.append(
                    f"WARNING {xlsx_path.name}: missing expected sheet: {sheet_name}"
                )
                continue

            ws = wb[sheet_name]
            hourly_by_day, records_by_day = read_hourly_by_day(ws)

            for count_date, hourly_totals in hourly_by_day.items():
                all_data[count_date][sheet_name] = dict(hourly_totals)
                records_by_station_day[count_date] += records_by_day[count_date]

        class_order = {name: index for index, name in enumerate(CLASS_SHEETS)}

        for count_date in sorted(all_data):
            for vehicle_class in sorted(
                all_data[count_date],
                key=lambda value: class_order.get(value, 999),
            ):
                hourly_totals = all_data[count_date][vehicle_class]

                row = {
                    "Station": base_station_id,
                    "CountDate": access_date(count_date),
                    "VehicleClass": vehicle_class,
                    "Total": sum(hourly_totals.values()),
                    "SourceFile": source_file,
                    "SourceSheet": vehicle_class,
                }

                for hour in range(1, 25):
                    row[f"H{hour:02d}"] = hourly_totals.get(hour, 0)

                class_rows.append(row)

            first_dt = datetime.combine(count_date, time.min)
            last_dt = first_dt + timedelta(days=1)

            station_rows.append(
                {
                    "County": DEFAULT_COUNTY,
                    "StationID": base_station_id,
                    "FirstDate": access_datetime(first_dt),
                    "LastDate": access_datetime(last_dt),
                    "RecordsLoaded": records_by_station_day[count_date],
                    "DaysCollected": 1,
                    "StationType": DEFAULT_STATION_TYPE,
                    "Route": DEFAULT_ROUTE,
                    "Description": description,
                    "CountyOrder": DEFAULT_COUNTY_ORDER,
                }
            )

        if not class_rows:
            warnings.append(f"WARNING {xlsx_path.name}: no class rows extracted")
            return class_rows, station_rows, warnings, success

        success = True
        return class_rows, station_rows, warnings, success

    except Exception as exc:
        warnings.append(f"SKIPPED {xlsx_path.name}: processing failed: {exc}")
        return [], [], warnings, False

    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    source_options = parser.add_mutually_exclusive_group()
    source_options.add_argument("--source", type=Path, help="Parent folder; all subfolders are searched")
    source_options.add_argument("--choose-folder", action="store_true", help="Open a folder picker")
    parser.add_argument("--output", type=Path, default=OUTPUT_FOLDER, help="Parent output folder")
    parser.add_argument("--pattern", default=INPUT_PATTERN, help="Filename filter (default: %(default)s)")
    args = parser.parse_args()
    source = args.source or SOURCE_FOLDER
    if args.choose_folder:
        source = choose_source(source)
        if source is None:
            print("Cancelled. No output written.")
            return 0
    source = source.expanduser().resolve()
    files = discover_workbooks(source, args.pattern)
    if not files:
        raise FileNotFoundError(
            f"No .xlsx files matching {args.pattern!r} found in {source} or its subfolders. "
            "Check the filename filter and workbook format. No output written."
        )

    output_parent = args.output.expanduser().resolve()
    output_parent.mkdir(parents=True, exist_ok=True)
    output = Path(tempfile.mkdtemp(
        prefix=datetime.now().strftime("atr_%Y%m%d_%H%M%S_"), dir=output_parent
    ))
    log_lines = []

    def report(message: str) -> None:
        print(message, flush=True)
        log_lines.append(message)

    report(f"Source: {source}")
    report("Searching all subfolders; reading original files directly.")
    report(f"Pattern: {args.pattern}")
    report(f"Matching workbooks: {len(files)}")
    report(f"Output: {output}")
    report("NOTE: Missing/invalid values and missing hours become zero, as in the original script.")
    report("NOTE: Totals combine all count columns; partial days are not excluded.")
    all_class_rows = []
    all_station_rows = []
    processed = 0
    warning_count = 0
    seen_station_days = {}
    for index, path in enumerate(files, 1):
        label = path.relative_to(source).as_posix()
        report(f"[{index}/{len(files)}] {label}")
        class_rows, station_rows, warnings, success = process_workbook(path, label)
        if success:
            processed += 1
            for row in station_rows:
                key = (row["StationID"], row["FirstDate"])
                if key in seen_station_days:
                    warnings.append(
                        f"POSSIBLE DUPLICATE station/day: {key}; "
                        f"{seen_station_days[key]} and {label}. Both retained; review before import."
                    )
                else:
                    seen_station_days[key] = label
            all_class_rows.extend(class_rows)
            all_station_rows.extend(station_rows)
        report(f"  Station-day rows: {len(station_rows)}; class rows: {len(class_rows)}")
        for warning in warnings:
            report(f"  {label}: {warning}")
        warning_count += len(warnings)

    class_fields = ["Station", "CountDate", "VehicleClass"] + [
        f"H{hour:02d}" for hour in range(1, 25)
    ] + ["Total", "SourceFile", "SourceSheet"]
    station_fields = [
        "County", "StationID", "FirstDate", "LastDate", "RecordsLoaded",
        "DaysCollected", "StationType", "Route", "Description", "CountyOrder",
    ]
    if all_class_rows:
        write_delimited(output / CLASS_OUTPUT_NAME, class_fields, all_class_rows)
        write_delimited(output / STATION_OUTPUT_NAME, station_fields, all_station_rows)
        report(f"Wrote: {output / CLASS_OUTPUT_NAME}")
        report(f"Wrote: {output / STATION_OUTPUT_NAME}")
    else:
        report("No usable class data found. No import tables written; see warnings.")
    report(f"Processed: {processed}/{len(files)}; warnings: {warning_count}")
    report(f"Class rows: {len(all_class_rows)}; station-day rows: {len(all_station_rows)}")
    if warning_count:
        report("Review the log before importing: files or class sheets may be missing, or counts duplicated.")
    report(f"Log: {output / LOG_OUTPUT_NAME}")
    write_log(output / LOG_OUTPUT_NAME, log_lines)
    return 0 if processed == len(files) and all_class_rows else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (OSError, RuntimeError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1) from None
