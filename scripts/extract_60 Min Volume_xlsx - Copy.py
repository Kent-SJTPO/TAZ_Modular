# extract_60_min_volume_total_to_access_txt.py
#
# Extracts total daily directional volume from Imperial ATR "60 Min Volume" workbooks.
#
# Source:
#   J:\Data Inventory\Traffic Counts\Imperial\02_May ATRs
#
# Processes files like:
#   ATR 12. Weekstown Rd, btwn Columbia Rd & Indian Cabin Rd_60 Min Volume.xlsx
#
# Ignores incompatible files like:
#   ATR 5. Deliah Road, near Wawa & Southern Driveway_60 Minute.xlsx
#
# Output:
#   C:\Users\Kschellinger\SJTPO_Git\TAZ_Modular\traffic_counts\output_csv
#
# Main output:
#   tblHourlyVolumeTotal_import.txt

from __future__ import annotations

import csv
import re
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SOURCE_FOLDER = Path(
    r"J:\Data Inventory\Traffic Counts\Imperial\02_May ATRs"
)

OUTPUT_FOLDER = Path(
    r"C:\Users\Kschellinger\SJTPO_Git\TAZ_Modular\traffic_counts\output_csv"
)

VOLUME_OUTPUT_NAME = "tblHourlyVolumeTotal_import.txt"
STATION_OUTPUT_NAME = "tblVolumeStations_import.txt"
LOG_OUTPUT_NAME = "volume_processing_log.txt"

SOURCE_SHEET = "Volume"
INFO_SHEET = "Count Information"


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def safe_int(value: Any) -> int:
    if value is None or value == "":
        return 0
    try:
        return int(value)
    except (TypeError, ValueError):
        pass
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def as_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if value is None or value == "":
        return None

    text = str(value).strip()

    formats = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%m/%d/%Y %I:%M:%S %p",
        "%m/%d/%Y %I:%M %p",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue

    return None


def access_date(value: date | None) -> str:
    if value is None:
        return ""
    return value.strftime("%m/%d/%Y")


def is_target_volume_file(path: Path) -> bool:
    name = path.name.lower()

    if not name.endswith(".xlsx"):
        return False

    if name.startswith("~$"):
        return False

    if "60 minute" in name:
        return False

    if "60 min volume" not in name:
        return False

    return True


def remove_atr_prefix(text: str) -> str:
    text = clean_text(text)

    text = re.sub(
        r"^\s*ATR\s*\d+\s*[\.\-_:]*\s*",
        "",
        text,
        flags=re.IGNORECASE,
    )

    return text.strip()


def read_count_information(wb: Any) -> dict[str, str]:
    info = {}

    if INFO_SHEET not in wb.sheetnames:
        return info

    ws = wb[INFO_SHEET]

    for row in ws.iter_rows(values_only=True):
        if not row:
            continue

        key = clean_text(row[0])

        if not key:
            continue

        value = ""

        for item in row[1:]:
            if clean_text(item):
                value = clean_text(item)
                break

        info[key.rstrip(":")] = value

    return info


def derive_station(info: dict[str, str], source_file: Path) -> str:
    count_name = clean_text(info.get("Count Name"))

    if count_name:
        return remove_atr_prefix(count_name)[:255]

    name = source_file.stem
    name = re.sub(r"_60\s*Min\s*Volume$", "", name, flags=re.IGNORECASE)

    return remove_atr_prefix(name)[:255]


def derive_route(info: dict[str, str], station: str) -> str:
    route = clean_text(info.get("Location 1"))

    if route:
        return route[:50]

    match = re.match(r"^\s*([^,]+)", station)

    if match:
        return match.group(1).strip()[:50]

    return ""


def parse_direction_header(header: str) -> tuple[str, str]:
    header = clean_text(header)

    if "," in header:
        direction, route = header.split(",", 1)
        return direction.strip(), route.strip()

    parts = header.split()

    if parts:
        return parts[0].strip(), " ".join(parts[1:]).strip()

    return header, ""


def find_header_row(ws: Any) -> int:
    max_check_row = min(ws.max_row, 25)

    for row_num in range(1, max_check_row + 1):
        first_cell = clean_text(ws.cell(row=row_num, column=1).value).lower()

        if first_cell in ("date/time", "datetime", "date time"):
            return row_num

    raise ValueError("Could not find Volume header row with Date/Time in column A")


def read_volume_rows(
    wb: Any,
    source_file: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[str]]:

    warnings = []
    volume_rows = []
    station_rows = []

    if SOURCE_SHEET not in wb.sheetnames:
        warnings.append(f"SKIPPED {source_file.name}: missing sheet '{SOURCE_SHEET}'")
        return volume_rows, station_rows, warnings

    info = read_count_information(wb)
    station = derive_station(info, source_file)
    default_route = derive_route(info, station)

    ws = wb[SOURCE_SHEET]
    header_row = find_header_row(ws)

    direction_columns = []

    for col in range(2, ws.max_column + 1):
        header = clean_text(ws.cell(row=header_row, column=col).value)

        if not header:
            continue

        direction, route_from_header = parse_direction_header(header)

        if not direction:
            continue

        direction_columns.append(
            {
                "column": col,
                "direction": direction,
                "route": route_from_header or default_route,
                "header": header,
            }
        )

    if not direction_columns:
        warnings.append(f"SKIPPED {source_file.name}: no direction columns found")
        return volume_rows, station_rows, warnings

    totals = defaultdict(lambda: defaultdict(int))
    record_counts = defaultdict(int)

    for row_num in range(header_row + 1, ws.max_row + 1):
        dt = as_datetime(ws.cell(row=row_num, column=1).value)

        if dt is None:
            continue

        count_date = dt.date()

        for item in direction_columns:
            col = item["column"]
            direction = item["direction"]
            value = safe_int(ws.cell(row=row_num, column=col).value)

            totals[count_date][direction] += value
            record_counts[count_date] += 1

    for count_date in sorted(totals):
        for item in direction_columns:
            direction = item["direction"]
            route = item["route"]

            volume_rows.append(
                {
                    "Station": station,
                    "CountDate": access_date(count_date),
                    "Direction": direction,
                    "Route": route,
                    "Total": totals[count_date].get(direction, 0),
                    "SourceFile": source_file.name,
                    "SourceSheet": SOURCE_SHEET,
                }
            )

        station_rows.append(
            {
                "Station": station,
                "CountDate": access_date(count_date),
                "Route": default_route,
                "Description": station,
                "SourceFile": source_file.name,
                "SourceSheet": INFO_SHEET,
                "RecordsLoaded": record_counts[count_date],
            }
        )

    return volume_rows, station_rows, warnings


def write_delimited(path: Path, fieldnames: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=fieldnames,
            delimiter="\t",
            quoting=csv.QUOTE_MINIMAL,
            lineterminator="\n",
            extrasaction="ignore",
        )

        writer.writeheader()

        for row in rows:
            writer.writerow(row)


def write_log(path: Path, lines: list[str]) -> None:
    with path.open("w", encoding="utf-8") as file:
        for line in lines:
            file.write(line + "\n")


def process_workbook(
    path: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[str], bool]:

    warnings = []
    wb = None

    try:
        wb = load_workbook(path, data_only=True, read_only=True)
        volume_rows, station_rows, read_warnings = read_volume_rows(wb, path)

        warnings.extend(read_warnings)

        if not volume_rows:
            warnings.append(f"WARNING {path.name}: no volume rows extracted")
            return volume_rows, station_rows, warnings, False

        return volume_rows, station_rows, warnings, True

    except Exception as exc:
        warnings.append(f"SKIPPED {path.name}: {exc}")
        return [], [], warnings, False

    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


def main() -> None:
    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)

    log_lines = []

    log_lines.append(f"Run time: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    log_lines.append(f"Source folder: {SOURCE_FOLDER}")
    log_lines.append(f"Output folder: {OUTPUT_FOLDER}")
    log_lines.append("Filter: *.xlsx files containing '60 Min Volume' and not '60 Minute'")
    log_lines.append("")

    print(f"Source folder: {SOURCE_FOLDER}")
    print(f"Output folder: {OUTPUT_FOLDER}")
    print("Filter: *.xlsx files containing '60 Min Volume' and not '60 Minute'")
    print("")

    if not SOURCE_FOLDER.exists():
        raise FileNotFoundError(f"Source folder not found: {SOURCE_FOLDER}")

    files = [
        path
        for path in sorted(SOURCE_FOLDER.glob("*.xlsx"))
        if is_target_volume_file(path)
    ]

    if not files:
        raise FileNotFoundError(
            f"No target '60 Min Volume' files found in {SOURCE_FOLDER}"
        )

    print(f"Files found: {len(files)}")
    log_lines.append(f"Files found: {len(files)}")

    for path in files:
        print(f"  {path.name}")
        log_lines.append(f"  {path.name}")

    print("")
    log_lines.append("")

    all_volume_rows = []
    all_station_rows = []
    all_warnings = []
    processed_count = 0

    for path in files:
        msg = f"Processing: {path.name}"
        print(msg)
        log_lines.append(msg)

        volume_rows, station_rows, warnings, success = process_workbook(path)

        all_volume_rows.extend(volume_rows)
        all_station_rows.extend(station_rows)
        all_warnings.extend(warnings)

        print(f"  Volume total rows: {len(volume_rows)}")
        print(f"  Station rows:      {len(station_rows)}")

        log_lines.append(f"  Volume total rows: {len(volume_rows)}")
        log_lines.append(f"  Station rows:      {len(station_rows)}")

        if success:
            processed_count += 1

        print("")
        log_lines.append("")

    volume_fields = [
        "Station",
        "CountDate",
        "Direction",
        "Route",
        "Total",
        "SourceFile",
        "SourceSheet",
    ]

    station_fields = [
        "Station",
        "CountDate",
        "Route",
        "Description",
        "SourceFile",
        "SourceSheet",
        "RecordsLoaded",
    ]

    volume_txt = OUTPUT_FOLDER / VOLUME_OUTPUT_NAME
    station_txt = OUTPUT_FOLDER / STATION_OUTPUT_NAME
    log_file = OUTPUT_FOLDER / LOG_OUTPUT_NAME

    write_delimited(volume_txt, volume_fields, all_volume_rows)
    write_delimited(station_txt, station_fields, all_station_rows)

    log_lines.append("Final summary")
    log_lines.append("------------------------------")
    log_lines.append(f"Files found:        {len(files)}")
    log_lines.append(f"Files processed:    {processed_count}")
    log_lines.append(f"Volume total rows:  {len(all_volume_rows)}")
    log_lines.append(f"Station rows:       {len(all_station_rows)}")
    log_lines.append(f"Volume TXT:         {volume_txt}")
    log_lines.append(f"Station TXT:        {station_txt}")

    if all_warnings:
        log_lines.append("")
        log_lines.append("Warnings")
        log_lines.append("------------------------------")
        for warning in all_warnings:
            log_lines.append(warning)

    write_log(log_file, log_lines)

    print("Wrote:")
    print(f"  {volume_txt}")
    print(f"  {station_txt}")
    print(f"  {log_file}")
    print("")
    print(f"Files processed:   {processed_count}")
    print(f"Volume total rows: {len(all_volume_rows)}")
    print(f"Station rows:      {len(all_station_rows)}")

    if all_warnings:
        print("")
        print("Warnings were written to the processing log.")


if __name__ == "__main__":
    main()