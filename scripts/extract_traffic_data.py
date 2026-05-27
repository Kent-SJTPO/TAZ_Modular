from pathlib import Path
import csv
from datetime import datetime, timedelta
from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Kschellinger\SJTPO_Git\TAZ_Modular")
INPUT_DIR = ROOT / "traffic_counts" / "input_xlsx"
OUTPUT_DIR = ROOT / "traffic_counts" / "output_csv"

VOLUME_OUT = OUTPUT_DIR / "njdot_atr_15min_volume_import.csv"
LENGTH_WIDE_OUT = OUTPUT_DIR / "njdot_atr_15min_length_wide_import.csv"
LENGTH_LONG_OUT = OUTPUT_DIR / "njdot_atr_15min_length_long_import.csv"
SUMMARY_OUT = OUTPUT_DIR / "njdot_traffic_import_summary.txt"


def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_datetime(value):
    if isinstance(value, datetime):
        return value

    text = clean_text(value)
    if not text:
        return None

    formats = [
        "%m/%d/%Y %I:%M %p",
        "%m/%d/%y %I:%M %p",
        "%Y-%m-%d %H:%M:%S",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass

    return None


def number_or_zero(value):
    if value is None or value == "":
        return 0
    try:
        return int(value)
    except Exception:
        try:
            return int(float(value))
        except Exception:
            return 0


def read_count_info(wb):
    info = {}

    if "Count Information" not in wb.sheetnames:
        return info

    ws = wb["Count Information"]

    for row in ws.iter_rows(values_only=True):
        key = clean_text(row[0]) if len(row) > 0 else ""
        value = clean_text(row[1]) if len(row) > 1 else ""

        if key:
            safe_key = (
                key.replace(" ", "")
                .replace("/", "")
                .replace("-", "")
                .replace("_", "")
            )
            info[safe_key] = value

    return info


def parse_direction_site(text):
    text = clean_text(text)

    if "," in text:
        direction, site = text.split(",", 1)
        return direction.strip(), site.strip()

    parts = text.split()
    if parts:
        return parts[0].strip(), text

    return "", text


def detect_workbook_type(wb):
    sheet_names = [s.lower().strip() for s in wb.sheetnames]

    if "volume" in sheet_names:
        return "volume"

    for ws in wb.worksheets:
        headers = [clean_text(c.value).lower() for c in ws[1]]
        if headers and headers[0] == "date/time":
            if any("to" in h or ">" in h for h in headers[1:]):
                return "length"

    return "unknown"


def process_volume(xlsx_path, wb, info):
    rows_out = []

    ws = wb["Volume"]
    headers = [clean_text(c.value) for c in ws[1]]

    count_name = info.get("CountName", xlsx_path.stem)
    site_code = info.get("SiteCode", "")
    station_id = info.get("StationID", "")

    for row in ws.iter_rows(min_row=2, values_only=True):
        dt = parse_datetime(row[0])
        if dt is None:
            continue

        end_dt = dt + timedelta(minutes=15)

        for col_index in range(1, len(headers)):
            direction, site_name = parse_direction_site(headers[col_index])
            volume = number_or_zero(row[col_index])

            rows_out.append({
                "SourceFile": xlsx_path.name,
                "CountName": count_name,
                "SiteName": site_name,
                "Direction": direction,
                "CountDate": dt.strftime("%Y-%m-%d"),
                "StartTime": dt.strftime("%H:%M:%S"),
                "EndTime": end_dt.strftime("%H:%M:%S"),
                "Volume": volume,
                "SiteCode": site_code,
                "StationID": station_id,
            })

    return rows_out


def safe_length_field(bin_name):
    text = clean_text(bin_name)
    text = text.replace(">", "GT")
    text = text.replace("<", "LT")
    text = text.replace(" ", "_")
    text = text.replace("-", "_")
    text = text.replace(".", "_")
    text = text.replace("to", "TO")
    return "Len_" + text


def process_length(xlsx_path, wb, info):
    wide_rows = []
    long_rows = []

    count_name = info.get("CountName", xlsx_path.stem)
    site_code = info.get("SiteCode", "")
    station_id = info.get("StationID", "")

    for ws in wb.worksheets:
        if ws.title.lower().strip() == "count information":
            continue

        headers = [clean_text(c.value) for c in ws[1]]
        if not headers or headers[0].lower() != "date/time":
            continue

        direction, site_name = parse_direction_site(ws.title)
        length_bins = headers[1:]

        for row in ws.iter_rows(min_row=2, values_only=True):
            dt = parse_datetime(row[0])
            if dt is None:
                continue

            end_dt = dt + timedelta(minutes=15)
            values = [number_or_zero(v) for v in row[1:1 + len(length_bins)]]
            interval_total = sum(values)

            base = {
                "SourceFile": xlsx_path.name,
                "CountName": count_name,
                "SiteName": site_name,
                "Direction": direction,
                "CountDate": dt.strftime("%Y-%m-%d"),
                "StartTime": dt.strftime("%H:%M:%S"),
                "EndTime": end_dt.strftime("%H:%M:%S"),
                "IntervalTotal": interval_total,
                "SiteCode": site_code,
                "StationID": station_id,
            }

            wide_row = dict(base)

            for bin_name, count_value in zip(length_bins, values):
                wide_row[safe_length_field(bin_name)] = count_value

                long_rows.append({
                    "SourceFile": xlsx_path.name,
                    "CountName": count_name,
                    "SiteName": site_name,
                    "Direction": direction,
                    "CountDate": dt.strftime("%Y-%m-%d"),
                    "StartTime": dt.strftime("%H:%M:%S"),
                    "EndTime": end_dt.strftime("%H:%M:%S"),
                    "LengthBin": bin_name,
                    "VehicleCount": count_value,
                    "SiteCode": site_code,
                    "StationID": station_id,
                })

            wide_rows.append(wide_row)

    return wide_rows, long_rows


def write_csv(path, rows):
    if not rows:
        return

    fieldnames = list(rows[0].keys())

    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    all_volume_rows = []
    all_length_wide_rows = []
    all_length_long_rows = []
    unknown_files = []

    files = sorted(INPUT_DIR.glob("*.xlsx"))

    for xlsx_path in files:
        if xlsx_path.name.startswith("~$"):
            continue

        wb = load_workbook(xlsx_path, data_only=True)
        info = read_count_info(wb)
        workbook_type = detect_workbook_type(wb)

        print(f"{xlsx_path.name}: {workbook_type}")

        if workbook_type == "volume":
            all_volume_rows.extend(process_volume(xlsx_path, wb, info))

        elif workbook_type == "length":
            wide_rows, long_rows = process_length(xlsx_path, wb, info)
            all_length_wide_rows.extend(wide_rows)
            all_length_long_rows.extend(long_rows)

        else:
            unknown_files.append(xlsx_path.name)

    write_csv(VOLUME_OUT, all_volume_rows)
    write_csv(LENGTH_WIDE_OUT, all_length_wide_rows)
    write_csv(LENGTH_LONG_OUT, all_length_long_rows)

    with SUMMARY_OUT.open("w", encoding="utf-8") as f:
        f.write("NJDOT traffic data import summary\n\n")
        f.write(f"Input folder: {INPUT_DIR}\n")
        f.write(f"Files scanned: {len(files)}\n")
        f.write(f"Volume records: {len(all_volume_rows)}\n")
        f.write(f"Length wide interval records: {len(all_length_wide_rows)}\n")
        f.write(f"Length long bin records: {len(all_length_long_rows)}\n")
        f.write(f"Volume total: {sum(r['Volume'] for r in all_volume_rows)}\n")
        f.write(f"Length total: {sum(r['IntervalTotal'] for r in all_length_wide_rows)}\n\n")
        f.write(f"Volume CSV: {VOLUME_OUT}\n")
        f.write(f"Length wide CSV: {LENGTH_WIDE_OUT}\n")
        f.write(f"Length long CSV: {LENGTH_LONG_OUT}\n\n")

        if unknown_files:
            f.write("Unknown files:\n")
            for name in unknown_files:
                f.write(f"- {name}\n")

    print("DONE")
    print(f"Files scanned: {len(files)}")
    print(f"Volume records: {len(all_volume_rows)}")
    print(f"Length wide records: {len(all_length_wide_rows)}")
    print(f"Length long records: {len(all_length_long_rows)}")
    print(f"Wrote summary: {SUMMARY_OUT}")


if __name__ == "__main__":
    main()