from pathlib import Path
import csv
from datetime import datetime, timedelta
from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Kschellinger\SJTPO_Git\TAZ_Modular")

INPUT_DIR = ROOT / "traffic_counts" / "input_xlsx"
OUTPUT_DIR = ROOT / "traffic_counts" / "output_csv"

OUT_CSV = OUTPUT_DIR / "njdot_atr_15min_volume_import.csv"
SUMMARY_OUT = OUTPUT_DIR / "njdot_atr_15min_volume_summary.txt"


def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_datetime(value):
    if isinstance(value, datetime):
        return value

    text = clean_text(value)
    if not text:
        return None

    formats = [
        "%m/%d/%Y %I:%M %p",
        "%m/%d/%y %I:%M %p",
        "%Y-%m-%d %H:%M:%S",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass

    raise ValueError(f"Could not parse date/time value: {text}")


def number_or_zero(value):
    if value is None or value == "":
        return 0

    try:
        return int(value)
    except Exception:
        return 0


def read_count_info(ws):
    info = {}

    for row in ws.iter_rows(values_only=True):
        key = clean_text(row[0]) if len(row) > 0 else ""
        value = clean_text(row[1]) if len(row) > 1 else ""

        if key:
            safe_key = (
                key.replace(" ", "")
                .replace("/", "")
                .replace("-", "")
                .replace("_", "")
            )
            info[safe_key] = value

    return info


def parse_direction_header(header):
    text = clean_text(header)

    if "," in text:
        direction, site_name = text.split(",", 1)
        return direction.strip(), site_name.strip()

    parts = text.split()
    if parts:
        return parts[0].strip(), text

    return "", text


def process_workbook(xlsx_path):
    wb = load_workbook(xlsx_path, data_only=True)

    count_info = {}
    if "Count Information" in wb.sheetnames:
        count_info = read_count_info(wb["Count Information"])

    count_name = count_info.get("CountName", xlsx_path.stem)
    start_date = count_info.get("StartDate", "")
    start_time = count_info.get("StartTime", "")
    site_code = count_info.get("SiteCode", "")
    station_id = count_info.get("StationID", "")

    if "Volume" not in wb.sheetnames:
        return []

    ws = wb["Volume"]

    headers = [clean_text(c.value) for c in ws[1]]

    rows_out = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        dt = parse_datetime(row[0])
        if dt is None:
            continue

        end_dt = dt + timedelta(minutes=15)

        for col_index in range(1, len(headers)):
            direction, site_name = parse_direction_header(headers[col_index])
            volume = number_or_zero(row[col_index])

            rows_out.append({
                "SourceFile": xlsx_path.name,
                "CountName": count_name,
                "SiteName": site_name,
                "Direction": direction,
                "CountDate": dt.strftime("%Y-%m-%d"),
                "StartTime": dt.strftime("%H:%M:%S"),
                "EndTime": end_dt.strftime("%H:%M:%S"),
                "Volume": volume,
                "StartDateFromInfo": start_date,
                "StartTimeFromInfo": start_time,
                "SiteCode": site_code,
                "StationID": station_id,
            })

    return rows_out


def write_csv(path, rows):
    if not rows:
        return

    fieldnames = list(rows[0].keys())

    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    all_rows = []

    files = sorted(INPUT_DIR.glob("*15 Min Volume*.xlsx"))

    for xlsx_path in files:
        if xlsx_path.name.startswith("~$"):
            continue

        rows = process_workbook(xlsx_path)
        all_rows.extend(rows)

    write_csv(OUT_CSV, all_rows)

    total_volume = sum(r["Volume"] for r in all_rows)

    with SUMMARY_OUT.open("w", encoding="utf-8") as f:
        f.write("NJDOT ATR 15-minute volume import summary\n")
        f.write("\n")
        f.write(f"Input folder: {INPUT_DIR}\n")
        f.write(f"Files processed: {len(files)}\n")
        f.write(f"Import records: {len(all_rows)}\n")
        f.write(f"Total volume: {total_volume}\n")
        f.write("\n")
        f.write(f"CSV: {OUT_CSV}\n")

    print("DONE")
    print(f"Files processed: {len(files)}")
    print(f"Import records: {len(all_rows)}")
    print(f"Total volume: {total_volume}")
    print(f"Wrote: {OUT_CSV}")
    print(f"Wrote: {SUMMARY_OUT}")


if __name__ == "__main__":
    main()