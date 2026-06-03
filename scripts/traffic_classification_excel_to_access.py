from pathlib import Path
import pandas as pd
from openpyxl import load_workbook


INPUT_FOLDER = Path(
    r"C:\Users\Kschellinger\SJTPO_Git\TAZ_Modular\traffic_counts\input_xlsx"
)

OUTPUT_CSV = Path(
    r"C:\Users\Kschellinger\SJTPO_Git\TAZ_Modular\traffic_counts\processed\Class_By_Hour_Classification_Counts.csv"
)

COUNTIES = [
    "Atlantic",
    "Burlington",
    "Camden",
    "Cape May",
    "Cumberland",
    "Gloucester",
    "Salem",
]

INCLUDE_TOTAL_CLASS = True

FIELDS = [
    "Station",
    "VehicleClass",
    *[f"H{i:02d}" for i in range(1, 25)],
    "Total",
    "SourceFile",
    "SourceSheet",
]


def clean_value(value):
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        return value if value else None

    return value


def normalize_label(value):
    value = clean_value(value)

    if value is None:
        return ""

    return (
        str(value)
        .strip()
        .lower()
        .replace(" ", "")
        .replace("_", "")
        .replace("-", "")
        .replace(":", "")
    )


def safe_number(value):
    value = clean_value(value)

    if value is None:
        return 0

    if isinstance(value, (int, float)):
        return value

    try:
        return float(str(value).replace(",", ""))
    except ValueError:
        return 0


def find_matching_county(file_name):
    name = file_name.lower()

    for county in COUNTIES:
        if county.lower() in name:
            return county

    return None


def is_class_by_hour_file(xlsx_path):
    name = xlsx_path.stem.lower()

    if xlsx_path.name.startswith("~$"):
        return False

    if xlsx_path.suffix.lower() != ".xlsx":
        return False

    if "class by hour" not in name:
        return False

    if find_matching_county(name) is None:
        return False

    return True


def find_station(ws):
    # Primary method: find a cell labeled Station or Station:
    # and return the first nonblank cell to the right.
    for row_num in range(1, min(ws.max_row, 40) + 1):
        for col_num in range(1, ws.max_column + 1):
            label = normalize_label(ws.cell(row=row_num, column=col_num).value)

            if label == "station":
                for look_col in range(col_num + 1, min(ws.max_column, col_num + 10) + 1):
                    possible = clean_value(ws.cell(row=row_num, column=look_col).value)

                    if possible is not None:
                        return str(possible).strip()

    # Fallback: common location in the files tested earlier.
    possible = clean_value(ws["L8"].value)

    if possible is not None:
        return str(possible).strip()

    return None


def find_header_row(ws):
    for row_num in range(1, min(ws.max_row, 80) + 1):
        row_values = [
            clean_value(ws.cell(row=row_num, column=col_num).value)
            for col_num in range(1, ws.max_column + 1)
        ]

        has_hour = any(
            str(value).strip().lower() == "hour"
            for value in row_values
            if value is not None
        )

        has_total = any(
            str(value).strip().lower() == "total"
            for value in row_values
            if value is not None
        )

        if has_hour and has_total:
            return row_num

    return None


def get_class_columns(ws, header_row):
    class_columns = []

    for col_num in range(1, ws.max_column + 1):
        header = clean_value(ws.cell(row=header_row, column=col_num).value)

        if header is None:
            continue

        header_text = str(header).strip()

        if header_text.lower() == "hour":
            continue

        if header_text.lower() == "total" and not INCLUDE_TOTAL_CLASS:
            continue

        class_columns.append((col_num, header_text))

    return class_columns


def process_sheet(ws, xlsx_path):
    output_rows = []

    station = find_station(ws)

    if not station:
        print(f"  Skipping sheet with no Station: {ws.title}")
        return output_rows

    header_row = find_header_row(ws)

    if header_row is None:
        print(f"  Skipping sheet with no Hour/Total header row: {ws.title}")
        return output_rows

    class_columns = get_class_columns(ws, header_row)

    if not class_columns:
        print(f"  Skipping sheet with no classification columns: {ws.title}")
        return output_rows

    first_hour_row = header_row + 1
    last_hour_row = first_hour_row + 23

    if last_hour_row > ws.max_row:
        print(f"  Skipping sheet with fewer than 24 hourly rows: {ws.title}")
        return output_rows

    for col_num, vehicle_class in class_columns:
        hourly_values = []

        for row_num in range(first_hour_row, last_hour_row + 1):
            hourly_values.append(safe_number(ws.cell(row=row_num, column=col_num).value))

        output_row = {
            "Station": station,
            "VehicleClass": vehicle_class,
            "Total": sum(hourly_values),
            "SourceFile": xlsx_path.name,
            "SourceSheet": ws.title,
        }

        for hour_num, value in enumerate(hourly_values, start=1):
            output_row[f"H{hour_num:02d}"] = value

        output_rows.append(output_row)

    return output_rows


def process_workbook(xlsx_path):
    wb = load_workbook(xlsx_path, data_only=True)
    output_rows = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        output_rows.extend(process_sheet(ws, xlsx_path))

    return output_rows


def main():
    print(f"Scanning folder: {INPUT_FOLDER}")

    if not INPUT_FOLDER.exists():
        raise FileNotFoundError(f"Input folder does not exist: {INPUT_FOLDER}")

    all_rows = []

    for xlsx_path in sorted(INPUT_FOLDER.glob("*.xlsx")):
        print(f"Found: {xlsx_path.name}")

        if not is_class_by_hour_file(xlsx_path):
            print(f"  Skipping: {xlsx_path.name}")
            continue

        print(f"  Processing: {xlsx_path.name}")
        rows = process_workbook(xlsx_path)
        print(f"  Records extracted: {len(rows)}")
        all_rows.extend(rows)

    df = pd.DataFrame(all_rows, columns=FIELDS)

    OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUTPUT_CSV, index=False)

    print("")
    print(f"Done. Wrote {len(df)} records.")
    print(f"Output CSV: {OUTPUT_CSV}")


if __name__ == "__main__":
    main()