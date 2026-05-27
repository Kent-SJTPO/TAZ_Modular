from pathlib import Path
import csv
from datetime import datetime, timedelta
from openpyxl import load_workbook


ROOT = Path(r"C:\Users\Kschellinger\SJTPO_Git\TAZ_Modular")

INPUT_DIR = ROOT / "traffic_counts" / "input_xlsx"
OUTPUT_DIR = ROOT / "traffic_counts" / "output_csv"

WIDE_OUT = OUTPUT_DIR / "njdot_atr_15min_wide_import.csv"
LONG_OUT = OUTPUT_DIR / "njdot_atr_15min_long_import.csv"
SUMMARY_OUT = OUTPUT_DIR / "njdot_atr_15min_import_summary.txt"


def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_datetime(value):
    if isinstance(value, datetime):
        return value

    text = clean_text(value)
    if not text:
        return None

    formats = [
        "%m/%d/%Y %I:%M %p",
        "%m/%d/%y %I:%M %p",
        "%Y-%m-%d %H:%M:%S",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass

    raise ValueError(f"Could not parse date/time value: {text}")


def read_count_info(ws):
    info = {}

    for row in ws.iter_rows(values_only=True):
        key = clean_text(row[0]) if len(row) > 0 else ""
        value = clean_text(row[1]) if len(row) > 1 else ""

        if key:
            safe_key = (
                key.replace(" ", "")
                .replace("/", "")
                .replace("-", "")
                .replace("_", "")
            )
            info[safe_key] = value

    return info


def direction_from_sheet_name(sheet_name):
    text = sheet_name.strip()

    if "," in text:
        return text.split(",", 1)[0].strip()

    first = text.split()[0].strip()
    return first


def site_from_sheet_name(sheet_name):
    text = sheet_name.strip()

    if "," in text:
        return text.split(",", 1)[1].strip()

    return text


def number_or_zero(value):
    if value is None or value == "":
        return 0

    try:
        return int(value)
    except Exception:
        return 0


def process_workbook(xlsx_path):
    wb = load_workbook(xlsx_path, data_only=True)

    count_info = {}
    if "Count Information" in wb.sheetnames:
        count_info = read_count_info(wb["Count Information"])

    count_name = count_info.get("CountName", xlsx_path.stem)
    start_date = count_info.get("StartDate", "")
    start_time = count_info.get("StartTime", "")
    site_code = count_info.get("SiteCode", "")
    station_id = count_info.get("StationID", "")

    wide_rows = []
    long_rows = []

    for ws in wb.worksheets:
        if ws.title.lower().strip() == "count information":
            continue

        direction = direction_from_sheet_name(ws.title)
        site_name = site_from_sheet_name(ws.title)

        headers = [clean_text(c.value) for c in ws[1]]
        if not headers or headers[0].lower() != "date/time":
            continue

        length_bins = headers[1:]

        for row in ws.iter_rows(min_row=2, values_only=True):
            dt = parse_datetime(row[0])
            if dt is None:
                continue

            end_dt = dt + timedelta(minutes=15)

            values = [number_or_zero(v) for v in row[1:1 + len(length_bins)]]
            interval_total = sum(values)

            base = {
                "SourceFile": xlsx_path.name,
                "CountName": count_name,
                "SiteName": site_name,
                "Direction": direction,
                "CountDate": dt.strftime("%Y-%m-%d"),
                "StartTime": dt.strftime("%H:%M:%S"),
                "EndTime": end_dt.strftime("%H:%M:%S"),
                "StartDateFromInfo": start_date,
                "StartTimeFromInfo": start_time,
                "SiteCode": site_code,
                "StationID": station_id,
                "IntervalTotal": interval_total,
            }

            wide_row = dict(base)
            for bin_name, count_value in zip(length_bins, values):
                field_name = "Len_" + bin_name.replace(">", "GT").replace(" ", "_").replace("to", "TO")
                wide_row[field_name] = count_value

                long_rows.append({
                    "SourceFile": xlsx_path.name,
                    "CountName": count_name,
                    "SiteName": site_name,
                    "Direction": direction,
                    "CountDate": dt.strftime("%Y-%m-%d"),
                    "StartTime": dt.strftime("%H:%M:%S"),
                    "EndTime": end_dt.strftime("%H:%M:%S"),
                    "LengthBin": bin_name,
                    "VehicleCount": count_value,
                    "SiteCode": site_code,
                    "StationID": station_id,
                })

            wide_rows.append(wide_row)

    return wide_rows, long_rows


def write_csv(path, rows):
    if not rows:
        return

    fieldnames = list(rows[0].keys())

    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def main():
    INPUT_DIR.mkdir(parents=True, exist_ok=True)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    all_wide_rows = []
    all_long_rows = []

    files = sorted(INPUT_DIR.glob("*.xlsx"))

    for xlsx_path in files:
        if xlsx_path.name.startswith("~$"):
            continue

        wide_rows, long_rows = process_workbook(xlsx_path)
        all_wide_rows.extend(wide_rows)
        all_long_rows.extend(long_rows)

    write_csv(WIDE_OUT, all_wide_rows)
    write_csv(LONG_OUT, all_long_rows)

    total_vehicles = sum(r["IntervalTotal"] for r in all_wide_rows)

    with SUMMARY_OUT.open("w", encoding="utf-8") as f:
        f.write("NJDOT ATR 15-minute import summary\n")
        f.write("\n")
        f.write(f"Input folder: {INPUT_DIR}\n")
        f.write(f"Files processed: {len(files)}\n")
        f.write(f"Wide interval records: {len(all_wide_rows)}\n")
        f.write(f"Long length-bin records: {len(all_long_rows)}\n")
        f.write(f"Total vehicles: {total_vehicles}\n")
        f.write("\n")
        f.write(f"Wide CSV: {WIDE_OUT}\n")
        f.write(f"Long CSV: {LONG_OUT}\n")

    print("DONE")
    print(f"Files processed: {len(files)}")
    print(f"Wide records: {len(all_wide_rows)}")
    print(f"Long records: {len(all_long_rows)}")
    print(f"Total vehicles: {total_vehicles}")
    print(f"Wrote: {WIDE_OUT}")
    print(f"Wrote: {LONG_OUT}")
    print(f"Wrote: {SUMMARY_OUT}")


if __name__ == "__main__":
    main()